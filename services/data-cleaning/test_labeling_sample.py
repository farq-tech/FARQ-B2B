"""Labelling sample on synthetic rows: python3 -m unittest test_labeling_sample (from services/data-cleaning)."""
import json
import pathlib
import subprocess
import sys
import tempfile
import unittest

import duckdb
from openpyxl import load_workbook

HERE = pathlib.Path(__file__).resolve().parent


class LabelingSampleTests(unittest.TestCase):
    def test_draws_only_unclassified_rows_half_with_a_raw_category(self):
        with tempfile.TemporaryDirectory() as d:
            d = pathlib.Path(d)
            db = duckdb.connect()
            rows = [(f'k{i}', ['none', 'directory_high'][i % 2 == 0 and i < 40], i == 1,
                     None if i % 3 else 'ورشة', f'name {i}', [None, 'Hotel', 'Unclassified', 'Restaurant'][i % 4])
                    for i in range(120)]
            db.execute('create table x(poi_key varchar, final_source varchar, invalid_name boolean, name_ar varchar, '
                       'name_en varchar, category varchar)')
            db.executemany('insert into x values (?,?,?,?,?,?)', rows)
            db.execute(f"copy (select poi_key, final_source, invalid_name from x) to '{d / 'm.parquet'}'")
            db.execute(f"""copy (select poi_key, name_ar, name_en, null::varchar source_license_name,
                null::varchar tradename, null::varchar description_ar, 'الرياض' city_ar, category from x)
                to '{d / 't.parquet'}'""")
            db.execute(f"copy (select poi_key, 'سيارات' cls_main_top from x) to '{d / 'e.parquet'}'")
            (d / 'tax.json').write_text(json.dumps({'subs': [{'sub_id': 'a.1', 'main_ar': 'سيارات'},
                                                             {'sub_id': 'b.1', 'main_ar': 'أغذية'}]}))
            out = d / 's.xlsx'
            cmd = [sys.executable, str(HERE / 'labeling_sample.py'), '--merged', str(d / 'm.parquet'),
                   '--targets', str(d / 't.parquet'), '--taxonomy', str(d / 'tax.json'),
                   '--estimates', str(d / 'e.parquet'), '--out', str(out), '--size', '40']
            subprocess.run(cmd, check=True, capture_output=True)
            ws = load_workbook(out)['sample']
            got = list(ws.iter_rows(min_row=2, values_only=True))
            keys = {r[0] for r in got}
            self.assertEqual(len(got), 40)
            self.assertNotIn('k1', keys, 'a non-descriptive name is not labelled')
            self.assertTrue(all(int(k[1:]) % 2 or int(k[1:]) >= 40 for k in keys), 'only rows still without a category')
            self.assertEqual(sum(r[7] is not None for r in got), 20)
            self.assertNotIn('Unclassified', {r[7] for r in got}, 'a placeholder category counts as none')
            self.assertEqual([c.value for c in load_workbook(out)['mains']['A']], ['سيارات', 'أغذية', 'لا يمكن التحديد'])
            r = subprocess.run(cmd, capture_output=True, text=True)
            self.assertNotEqual(r.returncode, 0, 'never overwrites')


if __name__ == '__main__':
    unittest.main()
